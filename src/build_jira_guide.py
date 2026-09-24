"""Build the step-by-step Jira Cloud Free guide (Vietnamese) for importing the IMS backlog."""
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ims_data import PROJECT, SPRINTS

BLUE = RGBColor(0x1F, 0x38, 0x64)


def set_font(run, size=12, bold=False, mono=False, color=None):
    name = "Consolas" if mono else "Times New Roman"
    run.font.name, run.font.size, run.bold = name, Pt(size), bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = color


def apply_widths(table, widths):
    """Set column widths in both tblGrid and every cell so Word and LibreOffice agree."""
    table.autofit = False
    grid = table._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths):
            gc.set(qn("w:w"), str(int(Cm(widths[i]).twips)))
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)


def shade(cell, hex_color):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


class Guide:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2), Cm(2), Cm(2)
        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = "Times New Roman", Pt(12)
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
        for lvl in (1, 2):
            st = self.doc.styles[f"Heading {lvl}"]
            st.font.name, st.font.color.rgb = "Times New Roman", BLUE
            st.font.size = Pt(16 if lvl == 1 else 13)
            st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")

    def h(self, text, lvl=1):
        self.doc.add_heading(text, lvl)

    def p(self, text, bold=False, italic=False):
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        set_font(run, bold=bold)
        run.italic = italic
        para.paragraph_format.space_after = Pt(4)
        return para

    def rich(self, parts):
        """parts: list of (text, bold) so UI labels can be bolded inline."""
        para = self.doc.add_paragraph()
        for text, bold in parts:
            set_font(para.add_run(text), bold=bold)
        para.paragraph_format.space_after = Pt(4)

    def steps(self, items):
        for i, item in enumerate(items, 1):
            para = self.doc.add_paragraph(style="List Number" if False else None)
            para.paragraph_format.left_indent = Cm(0.6)
            para.paragraph_format.first_line_indent = Cm(-0.6)
            para.paragraph_format.space_after = Pt(3)
            set_font(para.add_run(f"{i}. "), bold=True)
            parts = item if isinstance(item, list) else [(item, False)]
            for text, bold in parts:
                set_font(para.add_run(text), bold=bold)

    def bullets(self, items):
        for item in items:
            para = self.doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.6)
            para.paragraph_format.first_line_indent = Cm(-0.4)
            para.paragraph_format.space_after = Pt(2)
            set_font(para.add_run("• " + item))

    def code(self, text):
        table = self.doc.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        cell = table.rows[0].cells[0]
        shade(cell, "F2F2F2")
        cell.paragraphs[0].text = ""
        for i, line in enumerate(text.strip("\n").split("\n")):
            para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            set_font(para.add_run(line), size=10, mono=True)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def note(self, text, color="FFF2CC"):
        table = self.doc.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        cell = table.rows[0].cells[0]
        shade(cell, color)
        cell.paragraphs[0].text = ""
        set_font(cell.paragraphs[0].add_run(text), size=11)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def table(self, header, rows, widths=None):
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, text in enumerate(header):
            cell = t.rows[0].cells[i]
            shade(cell, "D9E2F3")
            cell.paragraphs[0].text = ""
            set_font(cell.paragraphs[0].add_run(text), size=11, bold=True)
        for r in rows:
            cells = t.add_row().cells
            for i, text in enumerate(r):
                cells[i].paragraphs[0].text = ""
                set_font(cells[i].paragraphs[0].add_run(str(text)), size=11)
        if widths:
            apply_widths(t, widths)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build(out):
    g = Guide()
    title = g.doc.add_paragraph()
    set_font(title.add_run("HƯỚNG DẪN TRIỂN KHAI DỰ ÁN IMS TRÊN JIRA CLOUD FREE"), size=18, bold=True, color=BLUE)
    g.p(f"Dự án: {PROJECT['name']}. Người thực hiện: {PROJECT['student']} (kiêm 5 vai trò PM, BA, FE, BE, TEST).")
    g.p("Mục tiêu: trong khoảng 20-30 phút, dựng xong project Jira Scrum với đầy đủ 10 Epic, 16 User Story, "
        "74 Task, 31 Sub-task và 6 Sprint, khớp 100% với file Gantt IMS_Gantt_Chart.xlsx.", italic=True)

    g.h("0. Các file đi kèm")
    g.table(["File", "Dùng để làm gì"], [
        ("IMS_Jira_Import.csv", "Backlog đầy đủ 131 work item để import (cột Sprint đang để trống, xem Bước 4)."),
        ("IMS_Jira_Test_5rows.csv", "File thử 5 dòng (1 Epic, 1 Story, 1 Task, 2 Sub-task) để chạy thử mapping trước."),
        ("src/build_jira_csv.py", "Script sinh lại CSV, điền Sprint ID và Assignee thật của bạn."),
        ("src/ims_data.py", "Nguồn dữ liệu duy nhất: sửa ở đây rồi chạy lại script, Gantt/CSV/báo cáo đều khớp."),
        ("IMS_Gantt_Chart.xlsx", "Baseline kế hoạch (Gantt) để đối chiếu với Jira Timeline."),
    ], widths=[5.5, 11])
    g.note("Quan trọng: KHÔNG mở rồi Save file CSV bằng Excel. Excel sẽ đổi mã hóa, làm hỏng tiếng Việt "
           "và đổi định dạng ngày. Nếu cần xem, mở bằng VS Code hoặc Notepad.")

    g.h("1. Cấu trúc backlog sẽ được tạo")
    g.table(["Loại work item", "Số lượng", "Nguồn trong kế hoạch", "Quan hệ cha"], [
        ("Epic", 10, "Nhóm WBS 1-10 (Quản trị dự án, Phân tích, Thiết kế, Hạ tầng, Sprint 1-5, Kiểm thử)", "-"),
        ("Story", 16, "US-01..US-16, có Story Point và Acceptance Criteria trong Description", "Epic Sprint tương ứng"),
        ("Task", 74, "Việc WBS không gắn User Story (PM, BA, TEST, hạ tầng) và 6 mốc M1-M6", "Epic nhóm WBS"),
        ("Sub-task", 31, "Việc BE/FE hiện thực một User Story, ví dụ 5.2 thuộc US-01", "Story"),
    ], widths=[2.8, 1.8, 8, 4])
    g.p("Mỗi work item có 3 label để lọc nhanh: vai trò (PM, BA, FE, BE, TEST), sprint (Sprint-0..Sprint-5, "
        "hoặc continuous cho việc lặp lại, milestone cho mốc) và mã WBS (WBS-5.2). Summary bắt đầu bằng [vai trò] "
        "và mã WBS, ví dụ \"[BE] 5.2 Viết API POST /auth/login...\", nên nhìn là biết đang đội mũ nào.")

    g.h("2. Tạo site Jira Cloud Free")
    g.steps([
        [("Mở ", False), ("https://www.atlassian.com/software/jira/free", True), (" và bấm ", False), ("Get it free", True), (".", False)],
        "Đăng ký bằng email (nên dùng email cá nhân bạn dùng lâu dài) và xác nhận qua hộp thư.",
        [("Đặt tên site, ví dụ ", False), ("tranvietanh-ims", True), (" → địa chỉ sẽ là tranvietanh-ims.atlassian.net.", False)],
        "Khi được hỏi mục đích, chọn Software development / Scrum; bỏ qua bước mời thành viên (bạn làm 1 mình, Free cho tối đa 10 user).",
    ])

    g.h("3. Tạo project Scrum loại company-managed")
    g.note("Bắt buộc chọn company-managed. Importer CSV của Jira chỉ giữ được quan hệ Epic → Story → Sub-task "
           "với project company-managed; project team-managed sẽ mất liên kết cha-con.", color="F8CBAD")
    g.steps([
        [("Thanh bên trái: ", False), ("Projects (Spaces) → Create project", True), (".", False)],
        [("Chọn template ", False), ("Scrum", True), (" → ", False), ("Use template", True), (".", False)],
        [("Ở màn hình chọn loại project, chọn ", False), ("Company-managed", True), (" (không chọn Team-managed).", False)],
        [("Name: ", False), ("IMS - Intern Management System", True), ("; Key: ", False), ("IM", True),
         (" (mọi work item sẽ có mã IM-1, IM-2...). Bấm ", False), ("Create", True), (".", False)],
        [("Bật ước lượng thời gian: ", False), ("Project settings → Features", True),
         (" (hoặc Board settings → Estimation): Estimation = Story points, bật Time tracking.", False)],
    ])

    g.h("4. Tạo 6 Sprint và lấy Sprint ID")
    g.p("Jira không tự tạo Sprint từ CSV và không nhận tên Sprint, chỉ nhận ID dạng số. Vì vậy phải tạo Sprint trước.")
    g.steps([
        [("Vào ", False), ("Backlog", True), (", bấm ", False), ("Create sprint", True), (" 6 lần.", False)],
        [("Đổi tên lần lượt (dấu ... → Edit sprint): ", False),
         (", ".join(s[1] for s in SPRINTS), True), (". Chưa cần bấm Start sprint.", False)],
        [("Lấy board ID: mở board, nhìn URL dạng .../boards/", False), ("1", True), (" → board ID = 1.", False)],
        [("Mở tab mới, dán địa chỉ sau (thay site và board ID):", False)],
    ])
    g.code("https://tranvietanh-ims.atlassian.net/rest/agile/1.0/board/1/sprint")
    g.p("Trình duyệt hiện JSON, mỗi sprint có \"id\" và \"name\". Ghi lại id theo thứ tự Sprint 0 → Sprint 5. "
        "Với site mới tạo thường là 1,2,3,4,5,6 nhưng vẫn phải kiểm tra, vì sprint mẫu Jira tự tạo có thể chiếm ID 1.")
    g.table(["Sprint", "Thời gian", "Mục tiêu", "ID của bạn"],
            [(s[1], f"{s[2]:%d/%m} - {s[3]:%d/%m/%Y}", s[5], "......") for s in SPRINTS], widths=[3, 3.6, 8, 2])

    g.h("5. Sinh CSV có Sprint ID và Assignee")
    g.p("Mở PowerShell trong thư mục dự án (C:\\Projects\\MyProject\\QLDAPM) và chạy, thay 6 số bằng ID thật "
        "và email bằng email tài khoản Atlassian của bạn:")
    g.code("cd C:\\Projects\\MyProject\\QLDAPM\\src\n"
           "python build_jira_csv.py ..\\out\\IMS_Jira_Import.csv --sprint-ids 1,2,3,4,5,6 --assignee email@cua-ban.com\n"
           "python build_jira_csv.py ..\\out\\IMS_Jira_Test_5rows.csv --sprint-ids 1,2,3,4,5,6 --assignee email@cua-ban.com --sample")
    g.p("Script in ra số dòng theo loại và dòng \"sprint column filled\". Khi đó: Story/Task nằm đúng Sprint theo ngày bắt đầu; "
        "Sub-task đi theo Story cha; 5 việc lặp lại (1.8, 1.9, 2.9, 9.4, 9.5) và các Epic để ở Backlog, không gán Sprint.")
    g.note("Không có Python? Vẫn import được file IMS_Jira_Import.csv có sẵn (cột Sprint trống), rồi xếp Sprint "
           "bằng bulk change ở Bước 8. Chậm hơn khoảng 5 phút.")

    g.h("6. Import thử 5 dòng (khuyến nghị)")
    g.p("Tạo thêm một project company-managed tạm tên IMSTEST, import IMS_Jira_Test_5rows.csv theo đúng Bước 7. "
        "Kiểm tra Story US-01 nằm dưới Epic 5, 2 Sub-task nằm dưới US-01, Story Point = 5, Sprint đúng. "
        "Nếu ổn, xóa project IMSTEST (Project settings → Details → ... → Move to trash) và import file đầy đủ vào IMS. "
        "Lưu lại file cấu hình ở cuối lần thử để lần import thật chỉ cần tải lên.")

    g.h("7. Import file CSV đầy đủ")
    g.p("7.1. Mở trình import", bold=True)
    g.steps([
        [("Bánh răng ", False), ("Settings (góc trên phải) → System", True), (". Cần quyền site admin, bạn là người tạo site nên đã có.", False)],
        [("Menu trái, mục Import and Export: ", False), ("External System Import", True), (" → chọn ", False), ("CSV", True), (".", False)],
        [("Lưu ý: không dùng \"Import work items from CSV\" trong màn hình Search/Filters, vì trình đó ", False),
         ("không import được quan hệ cha-con", True), (".", False)],
    ])
    g.p("7.2. Setup", bold=True)
    g.steps([
        [("Choose a CSV file: ", False), ("IMS_Jira_Import.csv", True), (". Nếu có file cấu hình từ lần thử, tick ", False),
         ("Use an existing configuration file", True), (".", False)],
        [("Mở ", False), ("Advanced", True), (": File encoding = ", False), ("UTF-8", True), (", CSV delimiter = ", False),
         (",", True), (" (dấu phẩy). Bấm Next.", False)],
        [("Import to Jira project: ", False), ("IM", True), ("; Date format: ", False), ("dd/MM/yyyy", True),
         (". Bấm Next.", False)],
    ])
    g.p("7.3. Map fields (bước quan trọng nhất)", bold=True)
    g.table(["Cột trong CSV", "Chọn Jira field", "Map field value?"], [
        ("Work item ID", "Work item Id (tên cũ: Issue Id)", "Không"),
        ("Work type", "Work type (tên cũ: Issue Type)", "Có"),
        ("Summary", "Summary", "Không"),
        ("Parent", "Parent (tên cũ: Parent Id)", "Không"),
        ("Description", "Description", "Không"),
        ("Priority", "Priority", "Có"),
        ("Assignee", "Assignee (bỏ qua nếu cột trống)", "Không"),
        ("Sprint", "Sprint", "Không"),
        ("Story point estimate", "Story point estimate (hoặc Story Points)", "Không"),
        ("Original Estimate", "Original Estimate (giá trị là giây: 14400 = 4h)", "Không"),
        ("Start date", "Start date (nếu không có trong danh sách thì bỏ qua)", "Không"),
        ("Due date", "Due date", "Không"),
        ("Labels (cả 3 cột)", "Labels", "Không"),
    ], widths=[4.2, 8.5, 3.3])
    g.p("7.4. Map values", bold=True)
    g.bullets([
        "Work type: Epic → Epic, Story → Story, Task → Task, Sub-task → Subtask (site mới gọi là \"Subtask\", không có gạch nối).",
        "Priority: Highest/High/Medium → giữ nguyên tên tương ứng.",
        "Bấm Validate: phải báo 131 work item, 0 lỗi. Có cảnh báo thì đọc kỹ (thường là tên trường Start date).",
        "Bấm Begin Import, chờ khoảng 1 phút, rồi bấm \"save the configuration\" để tải file cấu hình JSON về dùng lại.",
    ])

    g.h("8. Kiểm tra sau import (5 phút)")
    g.steps([
        "Backlog: IMS Sprint 0 có các Task chuẩn bị/thiết kế; IMS Sprint 1 có US-01..US-04 kèm Task BA/TEST. Tổng SP hiển thị ở đầu mỗi sprint: 13, 14, 16, 16.",
        "Mở US-01: có 2 Sub-task [BE] 5.2 và [FE] 5.3, Parent = Epic \"5. Sprint 1...\", Story point = 5, Description có 3 AC.",
        "Timeline: 10 Epic xếp đúng thời gian 31/08 → 20/11/2026, giống Gantt Excel.",
        "Tìm kiếm JQL để đếm: project = IM → 131 kết quả; project = IM AND issuetype = Subtask → 31.",
        [("Nếu đã import không có Sprint: Search với JQL ", False), ("project = IM AND labels = Sprint-1 AND issuetype in (Story, Task)", True),
         (" → ... → Bulk change → Edit fields → Sprint = IMS Sprint 1. Lặp lại cho Sprint-0..Sprint-5.", False)],
        [("Nếu chưa gán Assignee: JQL ", False), ("project = IM AND assignee is EMPTY", True),
         (" → Bulk change → Edit → Assignee = bạn.", False)],
    ])
    g.note("Import sai? Xóa hàng loạt bằng JQL project = IM → Bulk change → Delete, sửa file rồi import lại. "
           "Đừng import chồng lần 2 lên dữ liệu cũ vì sẽ bị nhân đôi.")

    g.h("9. Vận hành Scrum một người trên Jira")
    g.p("9.1. Cấu hình board", bold=True)
    g.bullets([
        "Board settings → Columns: To Do | In Progress | In Testing | Done. Thêm status IN TESTING để có bước TEST tách khỏi DEV.",
        "Board settings → Quick filters: tạo 5 filter labels = PM, labels = BA, labels = FE, labels = BE, labels = TEST. "
        "Đầu mỗi khung giờ bấm filter của vai trò đang đội mũ.",
        "Board settings → Card layout: hiện Story point estimate và Due date trên thẻ.",
        "WIP limit: cột In Progress tối đa 2 (giảm rủi ro R-01 quá tải khi 1 người làm 5 vai trò).",
    ])
    g.p("9.2. Nhịp mỗi Sprint (2 tuần)", bold=True)
    g.table(["Thời điểm", "Việc trên Jira", "Vai trò"], [
        ("Thứ Hai tuần 1", "Start sprint (đặt đúng ngày trong bảng Bước 4), ghi Sprint Goal", "PM"),
        ("Hằng ngày", "Kéo thẻ qua các cột, Log work giờ thực tế (dùng cho EVM - AC)", "Tất cả"),
        ("Khi xong DEV", "Chuyển Story sang In Testing, chạy test case, tạo Bug nếu lỗi", "TEST"),
        ("Thứ Năm tuần 2", "Nghiệm thu nội bộ theo AC, refine backlog Sprint sau", "BA"),
        ("Thứ Sáu tuần 2", "Complete sprint, chụp Burndown + Sprint report, viết Retrospective", "PM"),
    ], widths=[3.3, 10, 2.7])
    g.p("9.3. Ghi Bug", bold=True)
    g.bullets([
        "Create → Work type Bug; Summary \"[BUG] <mô tả ngắn>\"; Priority theo severity: S1 → Highest, S2 → High, S3 → Medium, S4 → Low.",
        "Description: môi trường, bước tái hiện, kết quả mong đợi/thực tế, ảnh chụp; Link \"is caused by\" tới Story/Sub-task liên quan.",
        "Label: bug-S1..bug-S4 và mã test case (TC-AUTH-02) để truy vết RTM.",
    ])
    g.p("9.4. Báo cáo cho môn học (Free plan có sẵn)", bold=True)
    g.bullets([
        "Reports → Burndown chart và Sprint report: chụp cuối mỗi Sprint làm evidence.",
        "Reports → Velocity chart: so sánh SP cam kết và hoàn thành (kế hoạch khoảng 15 SP/Sprint).",
        "Timeline: chụp so sánh với Gantt baseline để trình bày baseline và thực tế khi nghiệm thu.",
        "Kết nối GitHub: Apps → GitHub for Jira; đặt tên nhánh/commit chứa mã IM-xx để Jira tự gắn code vào work item.",
    ])

    g.h("10. Lỗi thường gặp và cách xử lý")
    g.table(["Hiện tượng", "Nguyên nhân", "Cách xử lý"], [
        ("Sub-task không nằm dưới Story", "Import vào project team-managed hoặc chưa map Parent", "Dùng project company-managed; map cột Parent → Parent"),
        ("Sprint trống sau import", "Cột Sprint để tên hoặc để trống", "Chạy lại script với --sprint-ids số; hoặc bulk change"),
        ("Tiếng Việt thành ký tự lạ", "File bị Excel lưu lại hoặc chọn sai encoding", "Sinh lại CSV bằng script, chọn UTF-8"),
        ("Ngày sai hoặc lỗi Date", "Date format trong wizard khác file", "Nhập đúng dd/MM/yyyy"),
        ("Estimate hiện 0m hoặc quá lớn", "Hiểu nhầm đơn vị", "Cột Original Estimate tính bằng giây (3600 = 1h)"),
        ("Không thấy Start date để map", "Trường chưa có trên màn hình project", "Bỏ qua cột, ngày vẫn có trong Due date/Description"),
        ("Không có \"Sub-task\" trong value mapping", "Site mới đổi tên thành \"Subtask\"", "Map Sub-task → Subtask"),
        ("Không thấy External System Import", "Tài khoản không phải site admin", "Đăng nhập tài khoản đã tạo site"),
    ], widths=[4.5, 5.5, 6])

    g.h("11. Tóm tắt nhanh (checklist 10 bước)")
    g.steps([
        "Tạo site Jira Free.",
        "Tạo project Scrum company-managed, key IM.",
        "Bật Story points và Time tracking.",
        "Tạo 6 sprint IMS Sprint 0..5.",
        "Lấy sprint ID qua /rest/agile/1.0/board/<id>/sprint.",
        "Chạy build_jira_csv.py với --sprint-ids và --assignee.",
        "Import thử file 5 dòng vào project tạm, lưu config.",
        "Import IMS_Jira_Import.csv vào IMS bằng External System Import (UTF-8, dd/MM/yyyy).",
        "Kiểm tra 131 item, 31 subtask, SP mỗi sprint 13/14/16/16.",
        "Cấu hình board (cột In Testing, quick filter theo vai trò, WIP 2) và Start Sprint 0.",
    ])
    g.h("Tài liệu tham khảo")
    g.bullets([
        "Atlassian - Import data from a CSV file: https://support.atlassian.com/jira-cloud-administration/docs/import-data-from-a-csv-file/",
        "Atlassian - Keep parent-child mapping during CSV import: https://support.atlassian.com/jira/kb/keep-issue-parent-child-mapping-during-csv-import-to-jira-cloud/",
        "Atlassian - Sprint ID must be a number: https://support.atlassian.com/jira/kb/jira-csv-import-sprint-id-must-be-number/",
        "Atlassian - Prepare a CSV file for import: https://support.atlassian.com/jira-software-cloud/docs/prepare-a-csv-file-for-import/",
    ])
    g.p("[retrieved 2026-09-24]. Giao diện Jira Cloud thay đổi thường xuyên (ví dụ Issue → Work item, Project → Space); "
        "nếu tên menu khác một chút, tìm theo chức năng tương đương.", italic=True)
    g.doc.save(out)
    print("saved", out)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "Huong_dan_Jira_IMS.docx")
