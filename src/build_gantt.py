"""Build the IMS Gantt chart workbook (same layout as the ITCare template)."""
import sys
from collections import defaultdict
from datetime import timedelta

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col

from ims_data import (CONTINUOUS_TASKS, CRITICAL_PATH, MILESTONE_GROUP, PROJECT, ROLE_NAMES,
                      SPRINTS, TEST_OWNER, USER_STORIES, WBS, is_group, workdays)

FIRST_DAY_COL = 9  # column I
HEADER_ROW = 10
COLORS = {
    "group_bar": "2F5597", "sub_bar": "9DC3E6", "test_bar": "A9D18E", "milestone": "C00000",
    "critical": "F4B183", "head": "DEEBF7", "day_head": "D6DCE4", "group_row": "D9E2F3",
    "due": "EAEEF3", "title": "1F3864",
}
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
FONT = "Calibri"


def fill(hex_color):
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


def working_days():
    days, cur = [], PROJECT["start"]
    while cur <= PROJECT["end"]:
        if cur.weekday() < 5:
            days.append(cur)
        cur += timedelta(days=1)
    return days


def style(cell, bold=False, size=10, color=None, bg=None, h="left", fmt=None, wrap=False):
    cell.font = Font(name=FONT, bold=bold, size=size, color=color)
    cell.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap)
    if bg:
        cell.fill = fill(bg)
    if fmt:
        cell.number_format = fmt


def build_gantt_sheet(ws, days):
    day_col = {day: FIRST_DAY_COL + i for i, day in enumerate(days)}
    last_col = FIRST_DAY_COL + len(days) - 1
    ws.sheet_view.showGridLines = False
    widths = {"A": 2.5, "B": 8, "C": 84, "D": 9, "E": 11, "F": 11, "G": 10, "H": 12}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    for c in range(FIRST_DAY_COL, last_col + 1):
        ws.column_dimensions[col(c)].width = 3.4

    ws["B1"] = "IMS - Gantt Chart (12 tuần / 6 Sprint)"
    style(ws["B1"], bold=True, size=22, color=COLORS["title"])
    ws.row_dimensions[1].height = 49.5
    info = [("Project Title", PROJECT["name"]),
            ("Project Manager", f"{PROJECT['student']} (kiêm PM, BA, FE, BE, TEST)"),
            ("Môn học", f"{PROJECT['course']} - {PROJECT['school']}"),
            ("Date", PROJECT["start"])]
    for i, (k, v) in enumerate(info, start=2):
        ws.merge_cells(f"B{i}:C{i}")
        ws.merge_cells(f"D{i}:H{i}")
        ws[f"B{i}"], ws[f"D{i}"] = k, v
        style(ws[f"B{i}"], bold=True, bg=COLORS["head"])
        style(ws[f"D{i}"], fmt="dd/mm/yyyy" if i == 5 else None)
        ws.row_dimensions[i].height = 24

    # Row 7: sprint bands, row 8: week bands, row 9: day letters
    for key, _name, s, e, title, *_ in SPRINTS:
        c1, c2 = day_col[s], day_col[e]
        ws.merge_cells(start_row=7, start_column=c1, end_row=7, end_column=c2)
        cell = ws.cell(7, c1, f"{title}\n{s:%d/%m/%Y} - {e:%d/%m/%Y}")
        style(cell, bold=True, size=9, bg=COLORS["head"], h="center", wrap=True)
    for w in range(0, len(days), 5):
        week = days[w:w + 5]
        c1 = day_col[week[0]]
        ws.merge_cells(start_row=8, start_column=c1, end_row=8, end_column=c1 + len(week) - 1)
        style(ws.cell(8, c1, f"WEEK {w // 5 + 1}\n{week[0]:%d/%m} - {week[-1]:%d/%m}"),
              bold=True, bg=COLORS["head"], h="center", wrap=True)
    heads = [("TASK", "ID"), ("TASK", "TITLE"), ("TASK", "OWNER"), ("START", "DATE"),
             ("DUE", "DATE"), ("DURATION", "IN DAYS"), ("PCT OF TASK", "COMPLETE")]
    for i, (a, b) in enumerate(heads):
        style(ws.cell(8, 2 + i, a), bold=True, size=9, bg=COLORS["day_head"], h="center")
        style(ws.cell(9, 2 + i, b), bold=True, size=9, bg=COLORS["day_head"], h="center")
    letters = "MTWRF"
    for day, c in day_col.items():
        label = f"{letters[day.weekday()]}\n{day.day}/{day.month}" if day.day == 1 or day == days[0] \
            else f"{letters[day.weekday()]}\n{day.day}"
        style(ws.cell(9, c, label), bold=True, size=8, bg=COLORS["day_head"], h="center", wrap=True)
    for r, h in ((7, 30), (8, 30), (9, 28.5)):
        ws.row_dimensions[r].height = h

    row = HEADER_ROW
    for tid, title, owner, start, due, _hours, _story in WBS:
        group = is_group(tid)
        milestone = tid.startswith(MILESTONE_GROUP + ".")
        values = [tid, title, owner, start, due,
                  f'=IF(OR(E{row}="",F{row}=""),"",NETWORKDAYS(E{row},F{row}))', 0]
        for i, v in enumerate(values):
            cell = ws.cell(row, 2 + i, v)
            style(cell, bold=group, bg=COLORS["group_row"] if group else (COLORS["due"] if i == 4 else None),
                  h="left" if i in (0, 1) else "center",
                  fmt={0: "@", 3: "dd/mm/yyyy", 4: "dd/mm/yyyy", 6: "0%"}.get(i))
            cell.border = BORDER
        if group:
            bar = COLORS["milestone"] if tid == MILESTONE_GROUP else COLORS["group_bar"]
        elif milestone:
            bar = COLORS["milestone"]
        elif tid in CRITICAL_PATH:
            bar = COLORS["critical"]
        elif owner == TEST_OWNER:
            bar = COLORS["test_bar"]
        else:
            bar = COLORS["sub_bar"]
        for day, c in day_col.items():
            cell = ws.cell(row, c)
            cell.border = BORDER
            in_range = start <= day <= due
            if group and tid == MILESTONE_GROUP:
                in_range = any(t[3] == day for t in WBS if t[0].startswith(MILESTONE_GROUP + "."))
            if in_range:
                cell.fill = fill(bar)
                if milestone:
                    cell.value = "◆"
                    style(cell, bold=True, size=9, color="FFFFFF", bg=bar, h="center")
        ws.row_dimensions[row].height = 19.5
        row += 1
    last_task_row = row - 1
    ws.conditional_formatting.add(f"H{HEADER_ROW}:H{last_task_row}",
                                  DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                                              color="5B9BD5"))
    ws.freeze_panes = ws.cell(HEADER_ROW, FIRST_DAY_COL)

    row += 1
    style(ws.cell(row, 2, "CHÚ GIẢI"), bold=True, size=12, color=COLORS["title"])
    legend = [
        (COLORS["group_bar"], "Công việc tổng hợp (Task/Epic) - bao trọn thời gian của các công việc con"),
        (COLORS["sub_bar"], "Công việc chi tiết do một vai trò PM, BA, FE hoặc BE thực hiện"),
        (COLORS["critical"], "Công việc nằm trên đường găng (Critical Path) - trễ 1 ngày là trễ cả dự án"),
        (COLORS["test_bar"], "Công việc kiểm thử do TEST phụ trách"),
        (COLORS["milestone"], "Mốc kiểm soát M1-M6 (◆), thời lượng 1 ngày, không tiêu tốn nguồn lực"),
    ]
    for color, text in legend:
        row += 1
        ws.cell(row, 2).fill = fill(color)
        style(ws.cell(row, 3, text))

    row += 2
    style(ws.cell(row, 2, "TÓM TẮT SPRINT"), bold=True, size=12, color=COLORS["title"])
    row += 1
    for i, h in enumerate(["Sprint", "Phạm vi chính", "Thời gian", "User Story", "", "", "SP"]):
        style(ws.cell(row, 2 + i, h), bold=True, bg=COLORS["head"], h="center" if i else "left")
    ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=7)
    sp_by_story = {u[0]: u[2] for u in USER_STORIES}
    total_sp = 0
    for idx, (key, name, s, e, _t, scope, stories) in enumerate(SPRINTS):
        row += 1
        sp = sum(sp_by_story[x] for x in stories)
        total_sp += sp
        vals = [f"Sprint {idx}", scope, f"Tuần {idx * 2 + 1}-{idx * 2 + 2}",
                f"{stories[0]}..{stories[-1]}" if stories else "-", None, None, sp]
        for i, v in enumerate(vals):
            if v is not None:
                style(ws.cell(row, 2 + i, v), h="left" if i < 2 else "center", wrap=i == 1)
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=7)
    row += 1
    for i, v in enumerate(["Tổng", "Release v1.0.0 - phạm vi FR-01..FR-30", "12 tuần",
                           f"{len(USER_STORIES)} User Story", None, None, total_sp]):
        if v is not None:
            style(ws.cell(row, 2 + i, v), bold=True, bg=COLORS["group_row"], h="left" if i < 2 else "center")
    ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=7)

    row += 2
    style(ws.cell(row, 2, "GHI CHÚ"), bold=True, size=12, color=COLORS["title"])
    total_hours = sum(t[5] for t in WBS if t[5])
    notes = [
        "Ô cần nhập khi triển khai: cột PCT OF TASK COMPLETE (H) - nhập % hoàn thành thực tế của từng dòng.",
        "Cột DURATION IN DAYS (G) là công thức NETWORKDAYS, tự tính số ngày làm việc từ START DATE và DUE DATE - không sửa tay.",
        f"Lịch tuần 5 ngày (Thứ Hai - Thứ Sáu), bắt đầu {PROJECT['start']:%d/%m/%Y}, kết thúc {PROJECT['end']:%d/%m/%Y}. "
        "Ngày lễ 02/09 được làm bù vào cuối tuần nên không trừ khỏi lịch.",
        "Toàn bộ 5 vai trò PM, BA, FE, BE, TEST do một người (Trần Việt Anh) đảm nhận; cột OWNER là vai trò đang 'đội mũ', "
        "không phải người khác nhau. Các dòng chồng thời gian được làm xen kẽ theo khung giờ trong ngày.",
        f"Tổng nỗ lực ước lượng {total_hours} giờ; tải trung bình khoảng 27 giờ/tuần, cao điểm 39 giờ ở tuần 2 "
        "(Sprint 0) - xem sheet 'Phân bổ nguồn lực'.",
        "Sprint 0 và Sprint 5 không cam kết Story Point vì dành cho chuẩn bị và ổn định; "
        f"{total_sp} SP chia cho 4 Sprint phát triển, velocity kế hoạch khoảng 15 SP/Sprint.",
        "Đường găng: " + " > ".join(CRITICAL_PATH) + ".",
        "Công việc " + ", ".join(CONTINUOUS_TASKS) + " chạy liên tục nhiều tuần theo nhịp Sprint, không phải làm một lần.",
        "Mỗi Sprint theo nhịp: BA chốt AC đầu Sprint > BE/FE phát triển > BA nghiệm thu nội bộ > TEST chạy test case > "
        "BA refine backlog cho Sprint kế tiếp.",
    ]
    for n in notes:
        row += 1
        style(ws.cell(row, 3, n), wrap=False)

    ws.print_title_rows = "7:9"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return last_task_row


def build_workload_sheet(ws, days):
    ws.sheet_view.showGridLines = False
    roles = list(ROLE_NAMES)
    load = defaultdict(float)
    for tid, _title, owner, start, due, hours, _s in WBS:
        if is_group(tid) or not hours:
            continue
        n = workdays(start, due)
        cur = start
        while cur <= due:
            if cur.weekday() < 5:
                week = (cur - PROJECT["start"]).days // 7 + 1
                load[(week, owner)] += hours / n
            cur += timedelta(days=1)
    ws["B1"] = "Phân bổ nguồn lực theo tuần (giờ) - 1 người, 5 vai trò"
    style(ws["B1"], bold=True, size=16, color=COLORS["title"])
    header = ["Tuần", "Từ ngày"] + roles + ["Tổng giờ", "Năng lực", "Đánh giá"]
    for i, h in enumerate(header):
        style(ws.cell(3, 2 + i, h), bold=True, bg=COLORS["head"], h="center")
        ws.column_dimensions[col(2 + i)].width = 13
    ws.column_dimensions[col(2 + len(header) - 1)].width = 34
    weeks = sorted({w for w, _ in load})
    for r, w in enumerate(weeks, start=4):
        style(ws.cell(r, 2, f"Tuần {w}"), h="center")
        style(ws.cell(r, 3, PROJECT["start"] + timedelta(days=7 * (w - 1))), h="center", fmt="dd/mm/yyyy")
        for i, role in enumerate(roles):
            style(ws.cell(r, 4 + i, round(load[(w, role)], 1)), h="center", fmt="0.0")
        tc = 4 + len(roles)
        first, last = col(4), col(tc - 1)
        style(ws.cell(r, tc, f"=SUM({first}{r}:{last}{r})"), bold=True, h="center", fmt="0.0")
        style(ws.cell(r, tc + 1, 40 if w <= 2 else 30), h="center")
        style(ws.cell(r, tc + 2, f'=IF({col(tc)}{r}>{col(tc + 1)}{r},"QUÁ TẢI - cần san tải","OK")'), h="center")
    end = 3 + len(weeks)
    r = end + 1
    style(ws.cell(r, 2, "Tổng"), bold=True, bg=COLORS["group_row"], h="center")
    for c in range(4, 5 + len(roles)):
        style(ws.cell(r, c, f"=SUM({col(c)}4:{col(c)}{end})"), bold=True, bg=COLORS["group_row"], h="center", fmt="0.0")
    ws.cell(r + 2, 2, "Năng lực: 40 giờ/tuần ở Sprint 0 (chưa vào kỳ thi), 30 giờ/tuần từ Sprint 1 do song song các môn học khác.")
    ws.cell(r + 3, 2, "Đã san tải (resource leveling): dời docker-compose sang tuần 1, khởi tạo React sang đầu Sprint 1, "
                      "giảm giờ các bản vẽ thiết kế; tuần 2 từ 56 giờ còn 39 giờ.")


def main(out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "IMS - Gantt Chart"
    days = working_days()
    build_gantt_sheet(ws, days)
    build_workload_sheet(wb.create_sheet("Phân bổ nguồn lực"), days)
    wb.save(out_path)
    print("saved", out_path)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "IMS_Gantt_Chart.xlsx")
