"""Build the full IMS project-management report (Vietnamese Word document)."""
import sys
from collections import defaultdict
from datetime import date, timedelta

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import ims_data as D

BLUE = RGBColor(0x1F, 0x38, 0x64)
FONT = "Times New Roman"
RATES = {"PM": 150_000, "BA": 120_000, "BE": 130_000, "FE": 120_000, "TEST": 100_000}
CONTINGENCY = 0.10
LEAVES = [t for t in D.WBS if not D.is_group(t[0]) and t[5]]


# ------------------------------------------------------------------ helpers
def set_font(run, size=13, bold=False, italic=False, mono=False, color=None):
    name = "Consolas" if mono else FONT
    run.font.name, run.font.size, run.bold, run.italic = name, Pt(size), bold, italic
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = color


def apply_widths(table, widths):
    """Set column widths in both tblGrid and every cell so Word and LibreOffice agree."""
    table.autofit = False
    grid = table._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths):
            gc.set(qn("w:w"), str(int(Cm(widths[i]).twips)))
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)


def shade(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    cell._element.get_or_add_tcPr().append(shd)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "separate"),
                      ("w:t", None), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if tag == "w:fldChar":
            el.set(qn("w:fldCharType"), attr)
        elif tag == "w:instrText":
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        else:
            el.text = "1" if "PAGE" in instr else "Mục lục sẽ hiện sau khi cập nhật trường."
        run._r.append(el)


def money(v):
    return f"{v:,.0f}".replace(",", ".") + " đ"


def spread(task):
    """Yield (day, hours) spreading a leaf task's hours evenly over its working days."""
    tid, _t, owner, start, due, hours, _s = task
    n = D.workdays(start, due)
    cur = start
    while cur <= due:
        if cur.weekday() < 5:
            yield cur, hours / n
        cur += timedelta(days=1)


class Report:
    def __init__(self):
        self.doc = Document()
        self.tables = defaultdict(int)
        self.chapter = 0
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin, sec.right_margin = Cm(2.5), Cm(2)
        sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = FONT, Pt(13)
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
        normal.paragraph_format.line_spacing = 1.3
        normal.paragraph_format.space_after = Pt(4)
        for lvl, size in ((1, 16), (2, 14), (3, 13)):
            st = self.doc.styles[f"Heading {lvl}"]
            st.font.name, st.font.size, st.font.bold = FONT, Pt(size), True
            st.font.color.rgb = BLUE
            rpr = st.element.get_or_add_rPr()
            fonts = rpr.get_or_add_rFonts()
            for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                fonts.set(qn(k), FONT)
            for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                fonts.attrib.pop(qn(k), None)
            st.paragraph_format.space_before = Pt(12 if lvl == 1 else 8)
            st.paragraph_format.space_after = Pt(6)

    # text blocks
    def h(self, text, lvl):
        if lvl == 1:
            self.chapter += 1
        self.doc.add_heading(text, lvl)

    def p(self, text, bold=False, italic=False, align=None, size=13):
        para = self.doc.add_paragraph()
        set_font(para.add_run(text), size=size, bold=bold, italic=italic)
        para.paragraph_format.first_line_indent = Cm(0 if align else 1)
        if align:
            para.alignment = align
        else:
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return para

    def bullets(self, items, size=13):
        for item in items:
            para = self.doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(1)
            para.paragraph_format.first_line_indent = Cm(-0.5)
            para.paragraph_format.space_after = Pt(2)
            if isinstance(item, tuple):
                set_font(para.add_run("- " + item[0] + ": "), size=size, bold=True)
                set_font(para.add_run(item[1]), size=size)
            else:
                set_font(para.add_run("- " + item), size=size)

    def mono(self, text):
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        cell = t.rows[0].cells[0]
        shade(cell, "F2F2F2")
        lines = text.strip("\n").split("\n")
        cell.paragraphs[0].text = ""
        for i, line in enumerate(lines):
            para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            para.paragraph_format.line_spacing = 1.0
            para.paragraph_format.space_after = Pt(0)
            set_font(para.add_run(line), size=9.5, mono=True)
        self.doc.add_paragraph()

    def table(self, caption, header, rows, widths=None, size=11, bold_rows=()):
        self.tables[self.chapter] += 1
        cap = self.doc.add_paragraph()
        set_font(cap.add_run(f"Bảng {self.chapter}.{self.tables[self.chapter]}: {caption}"), size=12,
                 bold=True, italic=True)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.keep_with_next = True
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = t.rows[0]
        tr_pr = hdr._tr.get_or_add_trPr()
        rep = OxmlElement("w:tblHeader")
        rep.set(qn("w:val"), "true")
        tr_pr.append(rep)
        for i, text in enumerate(header):
            cell = hdr.cells[i]
            shade(cell, "D9E2F3")
            cell.paragraphs[0].text = ""
            set_font(cell.paragraphs[0].add_run(str(text)), size=size, bold=True)
        for ri, r in enumerate(rows):
            cells = t.add_row().cells
            for i, text in enumerate(r):
                cells[i].paragraphs[0].text = ""
                lines = str(text).split("\n")
                for li, line in enumerate(lines):
                    para = cells[i].paragraphs[0] if li == 0 else cells[i].add_paragraph()
                    para.paragraph_format.line_spacing = 1.1
                    para.paragraph_format.space_after = Pt(0)
                    set_font(para.add_run(line), size=size, bold=ri in bold_rows)
                if ri in bold_rows:
                    shade(cells[i], "EAEEF3")
        if widths:
            apply_widths(t, widths)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------ numbers
def total_hours():
    return sum(t[5] for t in LEAVES)


def hours_by_role():
    out = defaultdict(float)
    for t in LEAVES:
        out[t[2]] += t[5]
    return out


def weekly_load():
    load = defaultdict(float)
    for t in LEAVES:
        for day, h in spread(t):
            load[((day - D.PROJECT["start"]).days // 7 + 1, t[2])] += h
    return load


def budget():
    by_role = hours_by_role()
    labour = {r: by_role[r] * RATES[r] for r in RATES}
    subtotal = sum(labour.values())
    return by_role, labour, subtotal, subtotal * CONTINGENCY, subtotal * (1 + CONTINGENCY)


def planned_value(until):
    return sum(h * RATES[t[2]] for t in LEAVES for day, h in spread(t) if day <= until)


# ------------------------------------------------------------------ sections
def cover(r):
    lines = [("HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG", 14, True),
             ("KHOA CÔNG NGHỆ THÔNG TIN 1", 14, True), ("", 12, False), ("", 12, False),
             ("BÁO CÁO BÀI TẬP LỚN", 20, True), ("MÔN QUẢN LÝ DỰ ÁN PHẦN MỀM", 18, True), ("", 12, False),
             ("ĐỀ TÀI:", 14, True), ("HỆ THỐNG QUẢN LÝ THỰC TẬP SINH", 18, True),
             ("(INTERN MANAGEMENT SYSTEM - IMS)", 14, True), ("", 12, False), ("", 12, False)]
    for text, size, bold in lines:
        para = r.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(para.add_run(text), size=size, bold=bold, color=BLUE if size >= 18 else None)
    info = [("Sinh viên thực hiện", D.PROJECT["student"]), ("Mã sinh viên", ".........."),
            ("Lớp", ".........."), ("Vai trò trong dự án", "PM, BA, FE, BE, TEST (một người đảm nhận)"),
            ("Giảng viên hướng dẫn", "..........")]
    for k, v in info:
        para = r.doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(3.5)
        set_font(para.add_run(f"{k}: "), bold=True)
        set_font(para.add_run(v))
    for _ in range(5):
        r.doc.add_paragraph()
    para = r.doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(para.add_run("Hà Nội, 2026"), bold=True, size=14)
    r.page_break()
    para = r.doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(para.add_run("MỤC LỤC"), size=16, bold=True, color=BLUE)
    add_field(r.doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    r.p("(Nhấn chuột phải vào mục lục > Update Field > Update entire table để cập nhật số trang.)",
        italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
    r.page_break()


def header_footer(r):
    sec = r.doc.sections[0]
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(hp.add_run("IMS - Intern Management System | Quản lý dự án phần mềm"), size=10, italic=True)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE")


def ch1_overview(r):
    r.h("1. Tổng quan dự án", 1)
    r.h("1.1. Bối cảnh và bài toán", 2)
    r.p("Mỗi kỳ, các doanh nghiệp phần mềm tiếp nhận hàng chục sinh viên thực tập. Hiện nay quy trình này "
        "thường được quản lý rời rạc bằng email, bảng tính và tin nhắn: hồ sơ đăng ký nằm trong hộp thư, danh sách "
        "Mentor nằm trong file Excel của phòng nhân sự, còn công việc, chấm công và báo cáo tuần mỗi Mentor theo dõi "
        "một kiểu. Hậu quả là bộ phận nhân sự khó nắm được tiến độ chung, Mentor mất thời gian tổng hợp, còn thực tập "
        "sinh không rõ mình đang làm tốt hay chưa.")
    r.p("Dự án IMS xây dựng một hệ thống web tập trung, quản lý thực tập sinh từ lúc đăng ký cho tới khi kết thúc "
        "kỳ thực tập, với 3 nhóm người dùng chính: Admin (nhân sự), Mentor và Intern. Đề tài có đủ các giai đoạn "
        "Requirement, Planning, Development, Testing, Bug Fix và Release nên phù hợp để áp dụng các lĩnh vực kiến "
        "thức quản lý dự án theo PMBOK: tích hợp, phạm vi, lịch trình, chi phí, chất lượng, nguồn lực, truyền thông, "
        "rủi ro và các bên liên quan.")
    r.h("1.2. Mục tiêu dự án (SMART)", 2)
    r.bullets([
        ("Cụ thể", "Xây dựng IMS phiên bản v1.0.0 đáp ứng 30 yêu cầu chức năng FR-01..FR-30 và 16 User Story."),
        ("Đo được", f"100% User Story đạt Acceptance Criteria; 100% test case mức Critical/High đạt; 0 bug S1, S2 tồn đọng khi phát hành."),
        ("Khả thi", f"Tổng nỗ lực ước lượng {total_hours()} giờ, trung bình khoảng 27 giờ/tuần cho một người."),
        ("Liên quan", "Số hóa toàn bộ quy trình thực tập, thay thế email và bảng tính rời rạc."),
        ("Có thời hạn", f"Hoàn thành trong 12 tuần, từ {D.PROJECT['start']:%d/%m/%Y} đến {D.PROJECT['end']:%d/%m/%Y}."),
    ])
    r.h("1.3. Phạm vi dự án", 2)
    r.p("Phạm vi được xác định trong Scope Statement (WBS 1.2) và là cơ sở để kiểm soát thay đổi. Mọi yêu cầu nằm "
        "ngoài danh sách dưới đây phải đi qua quy trình Change Request ở mục 6.10.")
    r.table("Phạm vi trong và ngoài dự án", ["Trong phạm vi (In-scope)", "Ngoài phạm vi (Out-of-scope)"], [
        ("Đăng nhập, phân quyền 3 role", "Tính lương, phụ cấp cho thực tập sinh"),
        ("Đăng ký thực tập trực tuyến, duyệt hồ sơ", "Tích hợp với hệ thống HRM/ERP có sẵn"),
        ("Quản lý hồ sơ Intern, Mentor, phòng ban", "Ứng dụng di động native (Android/iOS)"),
        ("Phân công Mentor, giao và theo dõi task", "Đăng nhập một lần (SSO), xác thực 2 lớp"),
        ("Chấm công, báo cáo tuần, đánh giá", "Chat thời gian thực, video call"),
        ("Theo dõi tiến độ, dashboard, xuất kết quả Excel/PDF", "Đa ngôn ngữ, triển khai lên cloud trả phí"),
    ], widths=[8.2, 8.2])
    r.h("1.4. Các bên liên quan (Stakeholder Register)", 2)
    r.table("Danh sách các bên liên quan", ["Bên liên quan", "Vai trò / mối quan tâm", "Mức ảnh hưởng", "Cách tương tác"], [
        ("Giảng viên hướng dẫn", "Sponsor học thuật, duyệt phạm vi và nghiệm thu", "Cao", "Báo cáo tiến độ theo mốc M1-M6"),
        ("Trần Việt Anh", "Đội dự án: PM, BA, FE, BE, TEST", "Cao", "Jira, GitHub, nhật ký hằng ngày"),
        ("Phòng nhân sự (Admin)", "Người dùng chính, cần theo dõi toàn bộ Intern", "Cao", "Phỏng vấn giả lập, UAT"),
        ("Mentor", "Giao việc, đánh giá; cần thao tác nhanh", "Trung bình", "Phỏng vấn, UAT"),
        ("Thực tập sinh", "Xem task, nộp báo cáo, xem đánh giá", "Trung bình", "Khảo sát, UAT"),
        ("Bạn cùng lớp", "Người dùng đại diện khi UAT, review chéo", "Thấp", "Buổi UAT tuần 12"),
    ], widths=[3.6, 6, 2.6, 4.2])
    r.h("1.5. Luồng nghiệp vụ tổng quát", 2)
    r.p("Luồng nghiệp vụ chính giữ nguyên như đề xuất ban đầu và được chi tiết hóa thành các Use Case ở chương 4:")
    r.mono("Ứng viên đăng ký (UC-04)\n   |\nAdmin duyệt hồ sơ (UC-05)  -->  Từ chối (kèm lý do)\n   |\n"
           "Phân phòng ban (UC-06)\n   |\nPhân Mentor (UC-07)\n   |\nMentor giao nhiệm vụ (UC-08)\n   |\n"
           "Intern thực hiện, chấm công (UC-08, UC-09)\n   |\nNộp báo cáo tuần (UC-10)\n   |\n"
           "Mentor đánh giá giữa kỳ / cuối kỳ (UC-11)\n   |\nXuất kết quả, hoàn thành kỳ thực tập (UC-12)")


def ch2_charter(r):
    by_role, labour, subtotal, reserve, total = budget()
    r.h("2. Project Charter", 1)
    r.p("Project Charter (WBS 1.1) là văn bản chính thức khởi động dự án, thuộc lĩnh vực Quản lý tích hợp. Charter "
        "trao quyền cho PM sử dụng nguồn lực và là căn cứ để đối chiếu khi nghiệm thu.")
    milestones = [t for t in D.WBS if t[0].startswith(D.MILESTONE_GROUP + ".")]
    r.table("Project Charter dự án IMS", ["Hạng mục", "Nội dung"], [
        ("Tên dự án", D.PROJECT["name"]),
        ("Người quản lý dự án", f"{D.PROJECT['student']} (kiêm BA, FE, BE, TEST)"),
        ("Nhà tài trợ", "Giảng viên môn Quản lý dự án phần mềm"),
        ("Mục đích", "Số hóa quy trình quản lý thực tập sinh từ đăng ký tới kết thúc kỳ thực tập."),
        ("Thời gian", f"{D.PROJECT['start']:%d/%m/%Y} - {D.PROJECT['end']:%d/%m/%Y} (12 tuần, 6 Sprint)"),
        ("Tiêu chí thành công",
         "SC-01: 16/16 User Story đạt Acceptance Criteria\n"
         "SC-02: 100% test case Critical/High đạt, 0 bug S1/S2 tồn đọng\n"
         "SC-03: Phát hành v1.0.0 đúng ngày 20/11/2026, lệch lịch mốc không quá 2 ngày\n"
         "SC-04: Chi phí thực tế không vượt ngân sách kể cả dự phòng\n"
         "SC-05: UAT Sign-off với 8/8 kịch bản đạt\n"
         "SC-06: Đủ hồ sơ: SRS, thiết kế, Test Report, User Manual, Release Notes"),
        ("Ngân sách (quy đổi công)", f"{money(total)} (gồm {int(CONTINGENCY * 100)}% dự phòng)"),
        ("Mốc chính", "\n".join(f"{t[1]}: {t[3]:%d/%m/%Y}" for t in milestones)),
        ("Giả định", "Làm việc khoảng 30 giờ/tuần; công cụ dùng gói miễn phí; yêu cầu ổn định sau M1."),
        ("Ràng buộc", "Một người thực hiện toàn bộ; thời hạn cố định theo lịch học kỳ; không có ngân sách tiền mặt."),
        ("Rủi ro chính", "Quá tải do kiêm nhiệm (R-01), thiếu kiểm thử độc lập (R-04), scope creep (R-03)."),
        ("Phê duyệt", "Giảng viên hướng dẫn: ..........          Người quản lý dự án: Trần Việt Anh"),
    ], widths=[4.2, 12.2])


def ch3_organization(r):
    r.h("3. Tổ chức dự án và quản lý nguồn lực", 1)
    r.h("3.1. Vai trò và trách nhiệm", 2)
    r.p("Dự án áp dụng mô hình Scrum với 5 vai trò chuyên môn. Điểm đặc biệt là cả 5 vai trò đều do một người, "
        "Trần Việt Anh, đảm nhận. Để vẫn giữ được tính kiểm soát, mỗi đầu việc trong WBS gắn đúng một vai trò chịu "
        "trách nhiệm (cột OWNER), và khi làm việc đó người thực hiện chỉ \"đội một chiếc mũ\": tư duy, tiêu chí và "
        "đầu ra theo đúng vai trò ấy.")
    r.table("Vai trò và công việc chính", ["Vai trò", "Trách nhiệm", "Sản phẩm bàn giao chính"], [
        ("PM - Project Manager", "Lập và theo dõi kế hoạch, Sprint, rủi ro, chi phí, báo cáo tiến độ",
         "Charter, Scope Statement, WBS, Gantt, Risk Register, Status Report, EVM"),
        ("BA - Business Analyst", "Khảo sát, đặc tả yêu cầu, viết User Story và AC, nghiệm thu nội bộ",
         "SRS (FR, NFR, BR), Use Case, Activity/State Diagram, User Story, RTM"),
        ("FE - Frontend Developer", "Thiết kế và hiện thực giao diện React theo wireframe",
         "Sitemap, wireframe Figma, 14 màn hình UI-01..UI-14"),
        ("BE - Backend Developer", "Kiến trúc, CSDL, API, bảo mật, triển khai",
         "ERD, OpenAPI, mã nguồn Spring Boot, docker-compose, Release Notes"),
        ("TEST - Tester", "Lập kế hoạch kiểm thử, viết và chạy test case, quản lý bug",
         "Test Plan, 40 test case, Postman/Selenium, Bug Report, Test Report"),
    ], widths=[3.8, 6.2, 6.4])
    r.h("3.2. Lịch luân phiên vai trò", 2)
    r.p("Để tránh chuyển ngữ cảnh liên tục, mỗi ngày được chia thành các khung giờ cố định (time-boxing). "
        "Khung giờ không dùng hết thì chuyển sang vai trò đang có việc trên đường găng.")
    r.table("Lịch đội mũ trong một ngày làm việc", ["Khung giờ", "Vai trò", "Việc điển hình"], [
        ("19:00 - 19:30", "PM", "Cập nhật Jira, burndown, log work, rà soát rủi ro"),
        ("19:30 - 20:30", "BA", "Chốt AC, trả lời câu hỏi nghiệp vụ, cập nhật RTM"),
        ("20:30 - 22:00", "BE / FE", "Phát triển theo Sub-task của Sprint (luân phiên theo ngày)"),
        ("22:00 - 22:30", "TEST", "Chạy test case cho item ở cột In Testing, ghi bug"),
        ("Thứ Bảy (4 giờ)", "BE / FE", "Khối việc lớn, tích hợp FE - BE"),
        ("Chủ nhật (1 giờ)", "PM", "Weekly Status Report, lập kế hoạch tuần tới"),
    ], widths=[3.5, 2.6, 10.3])
    r.h("3.3. Ma trận RACI", 2)
    r.p("Trong ma trận RACI, R (Responsible) là vai trò trực tiếp làm, A (Accountable) là vai trò chịu trách nhiệm "
        "cuối cùng, C (Consulted) được hỏi ý kiến, I (Informed) được thông báo. Tuy cùng một người, việc tách vai trò "
        "giúp xác định rõ tiêu chí kiểm tra: ví dụ code do BE viết (R) nhưng TEST phải xác nhận (C) và BA nghiệm thu "
        "theo AC trước khi PM đóng Story (A).")
    r.table("Ma trận RACI", ["Hoạt động", "PM", "BA", "FE", "BE", "TEST"], [
        ("Project Charter, kế hoạch dự án", "A/R", "C", "I", "C", "I"),
        ("Đặc tả yêu cầu, User Story, AC", "A", "R", "C", "C", "C"),
        ("Thiết kế giao diện", "I", "C", "A/R", "C", "I"),
        ("Thiết kế CSDL và API", "I", "C", "C", "A/R", "I"),
        ("Phát triển frontend", "I", "C", "A/R", "C", "I"),
        ("Phát triển backend", "I", "C", "C", "A/R", "I"),
        ("Viết và chạy test case", "I", "C", "C", "C", "A/R"),
        ("Nghiệm thu nội bộ theo AC", "A", "R", "C", "C", "C"),
        ("Quản lý rủi ro, thay đổi", "A/R", "C", "I", "C", "I"),
        ("Phát hành và bàn giao", "A", "C", "R", "R", "C"),
    ], widths=[6.4, 2, 2, 2, 2, 2])
    r.h("3.4. Biểu đồ phân bổ nguồn lực", 2)
    r.p("Giờ công của từng đầu việc được chia đều cho số ngày làm việc của nó để tính tải theo tuần. Năng lực là "
        "40 giờ/tuần trong Sprint 0 (chưa vào kỳ thi) và 30 giờ/tuần từ Sprint 1. Bản nháp đầu tiên cho tuần 2 lên tới "
        "56 giờ; sau khi san tải (resource leveling) bằng cách dời docker-compose sang tuần 1, dời khởi tạo React sang "
        "đầu Sprint 1 và giảm giờ các bản vẽ thiết kế, tuần cao điểm còn khoảng 39 giờ, nằm trong năng lực.")
    load = weekly_load()
    roles = list(D.ROLE_NAMES)
    rows, sums = [], defaultdict(float)
    for w in range(1, 13):
        vals = [load[(w, ro)] for ro in roles]
        cap = 40 if w <= 2 else 30
        for ro, v in zip(roles, vals):
            sums[ro] += v
        rows.append([f"Tuần {w}", (D.PROJECT["start"] + timedelta(days=7 * (w - 1))).strftime("%d/%m")]
                    + [f"{v:.1f}" for v in vals] + [f"{sum(vals):.1f}", cap, "OK" if sum(vals) <= cap else "Quá tải"])
    rows.append(["Tổng", ""] + [f"{sums[ro]:.0f}" for ro in roles] + [f"{sum(sums.values()):.0f}", "", ""])
    r.table("Tải công việc theo tuần (giờ)", ["Tuần", "Từ", *roles, "Tổng", "Năng lực", "Đánh giá"], rows,
            widths=[1.7, 1.5, 1.4, 1.4, 1.4, 1.4, 1.5, 1.6, 1.8, 2], size=10, bold_rows=(len(rows) - 1,))


def ch4_requirements(r):
    r.h("4. Phân tích yêu cầu", 1)
    r.p("Chương này là đầu ra của vai trò BA trong Sprint 0 (WBS nhóm 2), thuộc lĩnh vực Quản lý phạm vi. "
        "Yêu cầu được thu thập qua phỏng vấn giả lập 3 nhóm người dùng, sau đó đặc tả thành yêu cầu chức năng, "
        "phi chức năng, quy tắc nghiệp vụ, Use Case và User Story có Acceptance Criteria kiểm thử được.")
    r.h("4.1. Tác nhân của hệ thống", 2)
    r.table("Tác nhân", ["Tác nhân", "Mô tả"], D.ACTORS, widths=[4, 12.4])
    r.h("4.2. Yêu cầu chức năng", 2)
    r.table("Yêu cầu chức năng FR-01..FR-30", ["Mã", "Mô tả", "Tác nhân", "User Story"], D.FUNCTIONAL,
            widths=[1.8, 9.6, 2.8, 2.2], size=10.5)
    r.h("4.3. Yêu cầu phi chức năng", 2)
    r.table("Yêu cầu phi chức năng", ["Mã", "Nhóm", "Tiêu chí đo được"], D.NON_FUNCTIONAL, widths=[2, 3, 11.4])
    r.h("4.4. Quy tắc nghiệp vụ", 2)
    r.table("Business Rules", ["Mã", "Quy tắc"], D.BUSINESS_RULES, widths=[2, 14.4])
    r.h("4.5. Use Case", 2)
    r.p("Sơ đồ Use Case tổng quát gồm 4 tác nhân và 12 Use Case (vẽ bằng draw.io, đính kèm hồ sơ thiết kế). "
        "Ba Use Case quan trọng nhất được đặc tả chi tiết bên dưới.")
    r.table("Danh sách Use Case", ["Mã", "Tên Use Case", "Tác nhân"], D.USE_CASES, widths=[2, 8.4, 6])
    specs = [
        ("UC-04/UC-05: Đăng ký thực tập và duyệt hồ sơ",
         "Ứng viên mở trang đăng ký, nhập họ tên, email, trường, ngành, GPA, phòng ban mong muốn và tải CV PDF. "
         "Hệ thống kiểm tra BR-01, BR-02 rồi lưu đơn trạng thái PENDING. Admin mở danh sách đơn, xem CV, chọn Duyệt: "
         "hệ thống tạo tài khoản Intern, gửi email thông tin đăng nhập và chuyển đơn sang APPROVED.",
         "Email đã có đơn PENDING: báo lỗi, không lưu. CV sai định dạng hoặc quá 5 MB: báo lỗi. Admin chọn Từ chối: "
         "bắt buộc nhập lý do, đơn chuyển REJECTED và ứng viên nhận email."),
        ("UC-08: Giao và theo dõi task",
         "Mentor chọn Intern mình phụ trách, nhập tiêu đề, mô tả, độ ưu tiên, deadline. Hệ thống lưu task TODO và "
         "gửi thông báo. Intern chuyển task IN_PROGRESS, khi xong chuyển REVIEW. Mentor xem kết quả và chuyển DONE.",
         "Deadline ở quá khứ: báo lỗi. Mentor không phụ trách Intern: HTTP 403. Mentor chưa đạt: trả task về "
         "IN_PROGRESS kèm nhận xét. Intern sửa task người khác: HTTP 403."),
        ("UC-11: Đánh giá Intern",
         "Mentor mở hồ sơ Intern, chọn loại đánh giá giữa kỳ hoặc cuối kỳ, chấm 5 tiêu chí thang 10 (thái độ, kiến thức, "
         "kỹ năng, làm việc nhóm, kết quả công việc) và viết nhận xét. Hệ thống tính điểm trung bình, xếp loại theo "
         "BR-13 và thông báo cho Intern.",
         "Đã có đánh giá cùng loại: từ chối theo BR-12. Điểm ngoài 0..10: báo lỗi. Intern đã COMPLETED: chỉ xem."),
    ]
    for title, main, alt in specs:
        r.p(title, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        r.bullets([("Luồng chính", main), ("Luồng thay thế / ngoại lệ", alt)])
    r.h("4.6. Sơ đồ hoạt động và trạng thái", 2)
    r.p("Activity Diagram của luồng thực tập bám theo mục 1.5. Hai đối tượng có vòng đời quan trọng được mô tả "
        "bằng State Diagram như sau:")
    r.mono("HỒ SƠ INTERN\n  PENDING --duyệt--> APPROVED --tạo tài khoản, phân phòng ban--> ACTIVE\n"
           "  PENDING --từ chối--> REJECTED\n  ACTIVE --có đánh giá cuối kỳ + Admin kết thúc--> COMPLETED\n"
           "  ACTIVE --vi phạm / nghỉ giữa chừng--> TERMINATED\n\n"
           "TASK\n  TODO --Intern bắt đầu--> IN_PROGRESS --Intern nộp--> REVIEW --Mentor duyệt--> DONE\n"
           "                                   ^                        |\n"
           "                                   +------ Mentor trả lại --+")
    r.h("4.7. User Story và Acceptance Criteria", 2)
    r.p("Mỗi User Story viết theo mẫu \"Là <vai trò>, tôi muốn <mục tiêu> để <lợi ích>\" và có tối thiểu 2 "
        "Acceptance Criteria dạng Given - When - Then. Story Point được ước lượng theo dãy Fibonacci (1, 2, 3, 5, 8).")
    sprint_of_story = {us: s[1] for s in D.SPRINTS for us in s[6]}
    rows = [(us, f"{D.US_TITLES[us]}\n{story}", sp, prio, sprint_of_story[us].replace("IMS ", ""),
             "\n".join(f"AC{i}. {ac}" for i, ac in enumerate(acs, 1)))
            for us, _e, sp, prio, story, acs in D.USER_STORIES]
    r.table("Product Backlog: 16 User Story", ["Mã", "User Story", "SP", "Ưu tiên", "Sprint", "Acceptance Criteria"],
            rows, widths=[1.5, 4.6, 0.9, 1.6, 1.6, 6.2], size=9.5)


def ch5_design(r):
    r.h("5. Thiết kế hệ thống", 1)
    r.h("5.1. Kiến trúc tổng thể", 2)
    r.p("Hệ thống theo kiến trúc client - server: frontend là Single Page Application viết bằng React, giao tiếp "
        "với backend Spring Boot qua REST API bảo vệ bằng JWT. Backend chia 3 lớp Controller - Service - Repository; "
        "toàn bộ kiểm tra quyền sở hữu dữ liệu nằm ở lớp Service. Mọi thành phần chạy bằng Docker Compose để môi "
        "trường dev, test và UAT giống nhau.")
    r.mono("+---------------------+    HTTPS/JSON    +------------------------------+     JDBC     +-----------+\n"
           "|  React SPA (Vite)   | ---------------> |  Spring Boot REST API        | -----------> |  MySQL 8  |\n"
           "|  Tailwind, Recharts | <--------------- |  Controller > Service > Repo |              |  Flyway   |\n"
           "+---------------------+    JWT Bearer    |  Spring Security (JWT)       |              +-----------+\n"
           "                                         |  Apache POI / PDF, Mail      | --> ổ đĩa: CV, file xuất\n"
           "                                         +------------------------------+")
    r.table("Công nghệ sử dụng", ["Thành phần", "Công nghệ", "Lý do chọn"], [
        ("Frontend", "ReactJS 18, Vite, Tailwind CSS, React Router, Axios, Recharts", "Phổ biến, nhanh, nhiều tài liệu"),
        ("Backend", "Java 17, Spring Boot 3, Spring Security, Spring Data JPA", "Chuẩn doanh nghiệp, bảo mật tốt"),
        ("CSDL", "MySQL 8, Flyway migration", "Quen thuộc, quản lý phiên bản schema"),
        ("Triển khai", "Docker, Docker Compose", "Môi trường đồng nhất"),
        ("Quản lý mã nguồn", "Git, GitHub, GitHub Actions", "Miễn phí, CI tự động"),
        ("Quản lý dự án", "Jira Cloud Free (Scrum), Excel Gantt", "Miễn phí tới 10 người dùng"),
        ("Thiết kế", "Figma, draw.io", "Wireframe, UML"),
        ("Kiểm thử", "JUnit 5, Postman + Newman, Selenium WebDriver", "Unit, API, E2E"),
    ], widths=[3.2, 7.6, 5.6])
    r.h("5.2. Thiết kế cơ sở dữ liệu", 2)
    r.p(f"Cơ sở dữ liệu gồm {len(D.TABLES)} bảng, mở rộng từ 6 bảng của đề xuất ban đầu (USERS, DEPARTMENTS, INTERNS, "
        "TASKS, REPORTS, EVALUATIONS) để phục vụ đăng ký thực tập, Mentor, bình luận, chấm công, thông báo và nhật ký.")
    r.table("Danh sách bảng", ["Bảng", "Các cột chính"], D.TABLES, widths=[4, 12.4], size=10)
    r.h("5.3. Thiết kế API", 2)
    r.table(f"Danh sách {len(D.API)} endpoint REST", ["Method", "Endpoint", "Chức năng", "Quyền"], D.API,
            widths=[2, 6.4, 5, 3], size=9.5)
    r.p("Mã lỗi thống nhất: 400 dữ liệu không hợp lệ, 401 chưa đăng nhập hoặc token hết hạn, 403 không có quyền, "
        "404 không tìm thấy, 409 xung đột (ví dụ email trùng, đã chấm công).")
    r.h("5.4. Thiết kế giao diện", 2)
    r.table("Danh sách màn hình", ["Mã", "Màn hình", "Người dùng"], D.SCREENS, widths=[2, 9, 5.4])
    r.p("Ví dụ bố cục Intern Dashboard (UI-13, biến thể Intern):", align=WD_ALIGN_PARAGRAPH.LEFT)
    r.mono("--------------------------------------\n        INTERN DASHBOARD\n--------------------------------------\n"
           "Xin chào Nguyễn Văn A\nTask:\n[1] Thiết kế giao diện Login\n    Deadline: 28/08\n    Status: IN PROGRESS\n"
           "[2] Thiết kế Dashboard\n    Deadline: 02/09\n    Status: TODO\n--------------------------------------\n"
           "Tiến độ thực tập: 65%   [#############-------]\n"
           "Báo cáo tuần: đã nộp 5/6   Chấm công: 28/30 ngày")
    r.h("5.5. Cấu trúc mã nguồn", 2)
    r.mono("ims/\n  backend/   src/main/java/vn/ims/{auth,user,department,application,intern,mentor,\n"
           "             task,attendance,report,evaluation,dashboard,common}\n"
           "             src/main/resources/db/migration (Flyway V1__init.sql ...)\n"
           "  frontend/  src/{pages,components,services,hooks,routes}\n"
           "  tests/     postman/, selenium/\n  docs/      SRS, thiết kế, test report\n"
           "  docker-compose.yml   .github/workflows/ci.yml   README.md")


def ch6_plan(r):
    r.h("6. Kế hoạch dự án", 1)
    r.h("6.1. Phương pháp quản lý: Scrum", 2)
    r.p("Dự án chọn Scrum vì yêu cầu có thể được làm rõ dần, cần sản phẩm chạy được sau mỗi 2 tuần để giảng viên "
        "góp ý, và Jira hỗ trợ tốt quy trình này. Sprint 0 dành cho khởi động và thiết kế, 4 Sprint phát triển, "
        "Sprint 5 dành cho ổn định và phát hành. Các nghi thức Scrum được điều chỉnh cho đội một người:")
    r.bullets([
        ("Sprint Planning (1 giờ)", "PM chọn Story theo velocity, BA xác nhận AC, BE/FE tách Sub-task và ước lượng giờ."),
        ("Daily Scrum (10 phút)", "Viết nhật ký 3 câu hỏi vào comment Jira: hôm qua làm gì, hôm nay làm gì, vướng gì."),
        ("Sprint Review (1 giờ)", "Demo cho giảng viên hoặc bạn cùng lớp, ghi nhận phản hồi thành backlog item."),
        ("Retrospective (30 phút)", "Mẫu Start - Stop - Continue, ít nhất 1 hành động cải tiến cho Sprint sau."),
    ])
    r.p("Definition of Ready: Story có mô tả, tối thiểu 2 AC, đã ước lượng SP, không còn phụ thuộc chưa giải quyết. "
        "Definition of Done: code đã merge vào develop qua Pull Request tự review theo checklist, unit test tầng Service "
        "đạt, test case liên quan đạt, BA nghiệm thu theo AC, tài liệu API cập nhật.", align=None)
    r.h("6.2. Kế hoạch Sprint", 2)
    sp = {u[0]: u[2] for u in D.USER_STORIES}
    rows = [(s[1].replace("IMS ", ""), f"{s[2]:%d/%m} - {s[3]:%d/%m}", s[5],
             ", ".join(s[6]) if s[6] else "-", sum(sp[x] for x in s[6])) for s in D.SPRINTS]
    rows.append(("Tổng", "12 tuần", "Release v1.0.0", f"{len(D.USER_STORIES)} Story", sum(sp.values())))
    r.table("Kế hoạch 6 Sprint", ["Sprint", "Thời gian", "Mục tiêu", "User Story", "SP"], rows,
            widths=[2.2, 2.8, 6.6, 3.6, 1.2], bold_rows=(len(rows) - 1,))
    r.h("6.3. Cấu trúc phân rã công việc (WBS)", 2)
    r.p(f"WBS gồm 11 nhóm công việc, {len(LEAVES) + 6} đầu việc chi tiết (kể cả 6 mốc), được phân rã theo nguyên tắc "
        "100%: tổng các việc con bao trọn phạm vi của việc cha. Mỗi đầu việc có đúng một vai trò chịu trách nhiệm, "
        "thời lượng từ 1 đến 3 ngày (trừ các việc lặp lại theo Sprint) để dễ theo dõi.")
    rows, bold = [], []
    for tid, title, owner, start, due, hours, _s in D.WBS:
        if D.is_group(tid):
            bold.append(len(rows))
            hours = sum(t[5] for t in D.children(tid) if t[5])
        rows.append((tid, title, owner, f"{start:%d/%m}", f"{due:%d/%m}", D.workdays(start, due), hours or "-"))
    r.table("WBS chi tiết", ["ID", "Công việc", "Vai trò", "Bắt đầu", "Kết thúc", "Ngày", "Giờ"], rows,
            widths=[1.2, 9, 1.4, 1.5, 1.5, 1, 1], size=9, bold_rows=tuple(bold))
    r.h("6.4. Ước lượng", 2)
    total = total_hours()
    r.p("Ước lượng dùng hai cách bổ trợ. Thứ nhất, Story Point theo Planning Poker cho User Story: dùng US-04 "
        "(quản lý phòng ban, 2 SP) làm Story mốc, các Story khác so sánh tương đối. Thứ hai, ước lượng giờ công từ dưới "
        f"lên (bottom-up) cho từng đầu việc WBS, tổng cộng {total} giờ. Velocity kế hoạch khoảng 15 SP/Sprint; sau "
        "Sprint 1 sẽ đo velocity thực tế để hiệu chỉnh các Sprint sau.")
    pert = [(t[0], t[1], t[5] * 0.8, t[5], t[5] * 1.6) for t in LEAVES if t[0] in D.CRITICAL_PATH]
    te_sum = sum((o + 4 * m + p) / 6 for _i, _t, o, m, p in pert)
    sd = sum(((p - o) / 6) ** 2 for _i, _t, o, m, p in pert) ** 0.5
    r.p("Để kiểm tra độ tin cậy, các việc trên đường găng được ước lượng lại theo PERT ba điểm: lạc quan O = 0,8M, "
        "khả dĩ nhất M = giờ công baseline, bi quan P = 1,6M (bi quan lệch nhiều hơn vì một người làm, dễ bị gián đoạn). "
        "Công thức: TE = (O + 4M + P) / 6, độ lệch chuẩn SD = (P - O) / 6.")
    rows = [(i, t, f"{o:.1f}", f"{m:.0f}", f"{p:.1f}", f"{(o + 4 * m + p) / 6:.1f}") for i, t, o, m, p in pert]
    rows.append(("Tổng", "", f"{sum(x[2] for x in pert):.1f}", f"{sum(x[3] for x in pert):.0f}",
                 f"{sum(x[4] for x in pert):.1f}", f"{te_sum:.1f}"))
    r.table("Ước lượng PERT cho các việc trên đường găng (giờ)", ["ID", "Công việc", "O", "M", "P", "TE"], rows,
            widths=[1.2, 10, 1.3, 1.3, 1.3, 1.3], size=9.5, bold_rows=(len(rows) - 1,))
    base = sum(x[3] for x in pert)
    r.p(f"Kỳ vọng các việc găng là {te_sum:.1f} giờ so với baseline {base} giờ (cao hơn {te_sum / base * 100 - 100:.0f}%), "
        f"độ lệch chuẩn {sd:.1f} giờ. Với xác suất khoảng 95% (TE + 2SD), khối lượng găng không vượt "
        f"{te_sum + 2 * sd:.0f} giờ, tức phần tăng thêm nằm trong quỹ dự phòng 10% của cả dự án. Đây là căn cứ định "
        "lượng cho mức dự phòng ở mục 6.6.")
    r.h("6.5. Lịch trình, mốc kiểm soát và đường găng", 2)
    r.p("Lịch trình chi tiết được trình bày dưới dạng Gantt Chart trong file IMS_Gantt_Chart.xlsx (12 tuần, 6 Sprint), "
        "gồm cột ngày bắt đầu, kết thúc, số ngày làm việc (NETWORKDAYS) và % hoàn thành. Sheet 'Phân bổ nguồn lực' "
        "trình bày tải theo tuần ở mục 3.4. Các mốc kiểm soát là điểm đánh giá go/no-go với giảng viên:")
    ms = [t for t in D.WBS if t[0].startswith(D.MILESTONE_GROUP + ".")]
    crit = ["Scope, SRS, thiết kế được ký duyệt", "US-01..US-04 Done, 13 SP", "US-05..US-08 Done",
            "US-09..US-12 Done", "US-13..US-16 Done, đủ 30 FR", "UAT Sign-off, tag v1.0.0"]
    r.table("Mốc kiểm soát", ["Mốc", "Ngày", "Tiêu chí đạt"],
            [(t[1], f"{t[3]:%d/%m/%Y}", c) for t, c in zip(ms, crit)], widths=[8, 2.4, 6])
    titles = {t[0]: t for t in D.WBS}
    r.p("Đường găng (Critical Path) là chuỗi công việc dài nhất quyết định ngày kết thúc dự án, với độ trễ cho phép "
        "(float) bằng 0. Mọi việc trên đường găng được gắn Priority High trên Jira và tô màu cam trên Gantt:")
    r.table("Đường găng", ["ID", "Công việc", "Vai trò", "Thời gian"],
            [(i, titles[i][1], titles[i][2], f"{titles[i][3]:%d/%m} - {titles[i][4]:%d/%m}") for i in D.CRITICAL_PATH],
            widths=[1.3, 10.5, 1.6, 3], size=10)
    r.h("6.6. Ngân sách", 2)
    by_role, labour, subtotal, reserve, grand = budget()
    r.p("Dự án không phát sinh chi phí tiền mặt vì mọi công cụ dùng gói miễn phí. Để phục vụ quản lý chi phí và "
        "tính EVM, công sức được quy đổi thành tiền theo đơn giá tham khảo thị trường cho từng vai trò mức fresher.")
    rows = [(f"{ro} - {D.ROLE_NAMES[ro]}", f"{by_role[ro]:.0f}", money(RATES[ro]), money(labour[ro])) for ro in RATES]
    rows += [("Công cụ (Jira Free, GitHub, Figma, Postman, Docker)", "-", "-", money(0)),
             ("Cộng chi phí nhân công", f"{total:.0f}", "", money(subtotal)),
             (f"Dự phòng rủi ro {int(CONTINGENCY * 100)}% (Contingency reserve)", "", "", money(reserve)),
             ("Tổng ngân sách (BAC + dự phòng)", "", "", money(grand))]
    r.table("Ngân sách dự án (quy đổi)", ["Hạng mục", "Giờ", "Đơn giá/giờ", "Thành tiền"], rows,
            widths=[7.4, 1.8, 3.2, 4], bold_rows=(len(rows) - 3, len(rows) - 1))
    r.h("6.7. Kế hoạch truyền thông", 2)
    r.table("Kế hoạch truyền thông", ["Thông tin", "Người nhận", "Kênh", "Tần suất", "Người gửi"], [
        ("Nhật ký Daily Scrum", "Chính mình (lưu vết)", "Comment Jira", "Hằng ngày", "PM"),
        ("Weekly Status Report", "Giảng viên", "Email / Google Drive", "Chủ nhật hằng tuần", "PM"),
        ("Sprint Review, demo", "Giảng viên, bạn cùng lớp", "Gặp trực tiếp / Google Meet", "Cuối mỗi Sprint", "PM"),
        ("Burndown, Velocity", "Giảng viên", "Ảnh chụp Jira Reports", "Cuối mỗi Sprint", "PM"),
        ("Change Request", "Giảng viên", "Email + Jira", "Khi phát sinh", "PM"),
        ("Bug Report", "BE/FE (chính mình)", "Jira Bug", "Khi phát hiện", "TEST"),
        ("Báo cáo nghiệm thu", "Giảng viên", "Buổi bảo vệ", "Tuần 12", "PM"),
    ], widths=[3.8, 3.4, 3.8, 3, 2.4])
    r.h("6.8. Kế hoạch quản lý chất lượng", 2)
    r.p("Chất lượng được đảm bảo theo hai hướng: phòng ngừa (QA) và kiểm tra (QC). QA gồm chuẩn code, checklist review "
        "Pull Request, viết test case từ AC trước khi code, CI tự chạy build và unit test. QC gồm kiểm thử API, giao "
        "diện, tích hợp, hồi quy và UAT như chương 8. Chỉ số chất lượng theo dõi: tỉ lệ test case đạt, số bug theo "
        "mức độ, bug tái mở, độ phủ unit test tầng Service tối thiểu 60% (NFR-09).")
    r.h("6.9. Quản lý cấu hình", 2)
    r.bullets([
        ("Nhánh", "main (bản phát hành), develop (tích hợp), feature/IM-xx-mo-ta (mỗi Story/Sub-task), hotfix/IM-xx."),
        ("Commit", "\"IM-12: them API check-in\" để Jira tự liên kết commit với work item."),
        ("Pull Request", "Tự review theo checklist: đúng AC, có test, không lộ mật khẩu/khóa, không còn code thừa, CI xanh."),
        ("Phiên bản", "Tag cuối mỗi Sprint v0.1..v0.4, bản phát hành v1.0.0; migration CSDL đánh số Flyway."),
        ("Tài liệu", "SRS, thiết kế, test report lưu trong thư mục docs/ cùng repo để quản lý phiên bản."),
    ])
    r.h("6.10. Quản lý thay đổi", 2)
    r.p("Mọi thay đổi phạm vi, lịch hoặc chi phí sau mốc M1 phải lập Change Request (CR) gồm: mã CR, người đề xuất, "
        "mô tả, lý do, tác động phạm vi/lịch/chi phí/rủi ro, phương án, quyết định. Quy trình: ghi nhận CR trên Jira, "
        "BA phân tích tác động, PM đánh giá với baseline, xin ý kiến giảng viên nếu ảnh hưởng mốc, cập nhật backlog, "
        "Gantt và RTM nếu được duyệt. Thay đổi nhỏ trong một Story không làm đổi SP được PM tự quyết.")


def ch7_risk(r):
    r.h("7. Quản lý rủi ro", 1)
    r.p("Rủi ro được nhận diện ngay ở Sprint 0 (WBS 1.7) và rà soát lại mỗi tuần trong Weekly Status Report. "
        "Xác suất và tác động chấm theo thang 3 mức: Thấp = 1, Trung bình = 2, Cao = 3; điểm rủi ro = xác suất x "
        "tác động. Rủi ro từ 6 điểm trở lên là ưu tiên cao, được theo dõi hằng tuần.")
    score = {"Thấp": 1, "Trung bình": 2, "Cao": 3}
    rows = sorted(D.RISKS, key=lambda x: -score[x[2]] * score[x[3]])
    r.table("Risk Register", ["Mã", "Rủi ro", "XS", "TĐ", "Điểm", "Biện pháp ứng phó", "Chủ"],
            [(x[0], x[1], x[2], x[3], score[x[2]] * score[x[3]], x[4], x[5]) for x in rows],
            widths=[1.3, 4.6, 1.5, 1.5, 1.1, 5.4, 1.2], size=9.5)
    grid = {(p, i): [] for p in score for i in score}
    for x in D.RISKS:
        grid[(x[2], x[3])].append(x[0])
    r.table("Ma trận xác suất - tác động", ["Xác suất \\ Tác động", "Thấp", "Trung bình", "Cao"],
            [(p, *[", ".join(grid[(p, i)]) or "-" for i in ("Thấp", "Trung bình", "Cao")])
             for p in ("Cao", "Trung bình", "Thấp")], widths=[4, 4, 4, 4])
    r.p("Chiến lược ứng phó: giảm thiểu (Mitigate) cho R-01, R-04, R-05, R-06, R-07, R-10, R-11; tránh (Avoid) cho R-03 "
        "bằng quy trình Change Request; chuyển giao một phần (Transfer) cho R-04 bằng cách nhờ bạn cùng lớp làm UAT; "
        "chấp nhận (Accept) cho R-09 vì ảnh hưởng thấp. Quỹ dự phòng 10% ngân sách dùng cho rủi ro đã nhận diện; "
        "rủi ro R-01, R-02 là đặc thù của đội một người và được theo dõi qua biểu đồ tải tuần.")


def ch8_test(r):
    r.h("8. Kế hoạch kiểm thử", 1)
    r.p("Kế hoạch kiểm thử (WBS 9.1) xác định phạm vi, cấp độ, công cụ và tiêu chí đánh giá. Vì cùng một người vừa "
        "viết code vừa kiểm thử, test case được viết từ Acceptance Criteria trước khi code (nguyên tắc test-first) và "
        "ưu tiên kiểm thử tự động để giảm thiên kiến.")
    r.table("Cấp độ kiểm thử", ["Cấp độ", "Nội dung", "Công cụ", "Thời điểm"], [
        ("Unit test", "Tầng Service backend, độ phủ tối thiểu 60%", "JUnit 5, Mockito", "Trong mỗi Sub-task"),
        ("API test", "Kiểm tra endpoint, mã lỗi, phân quyền", "Postman, Newman", "Khi Story vào In Testing"),
        ("Integration test", "FE - BE - MySQL - lưu file", "Thủ công, Docker Compose", "Sprint 4"),
        ("System / E2E", "5 luồng chính trên trình duyệt", "Selenium WebDriver", "Sprint 4-5"),
        ("Regression", "Toàn bộ luồng Critical/High", "Newman + Selenium", "Tuần 11"),
        ("UAT", "8 kịch bản nghiệp vụ", "Người dùng đại diện", "Tuần 12"),
    ], widths=[3, 5.4, 4, 4])
    r.bullets([
        ("Tiêu chí bắt đầu", "Story đạt Definition of Done phía DEV, build CI xanh, dữ liệu test sẵn sàng."),
        ("Tiêu chí kết thúc", "100% test case Critical/High đạt, 0 bug S1/S2 mở, tối đa 3 bug S3 có phương án xử lý."),
    ])
    r.table("Mức độ nghiêm trọng của bug", ["Mức", "Định nghĩa", "Priority Jira", "Thời hạn sửa"], [
        ("S1 - Critical", "Sập hệ thống, mất dữ liệu, lộ dữ liệu người khác", "Highest", "Trong ngày"),
        ("S2 - Major", "Chức năng chính không dùng được, không có cách vòng", "High", "Trong Sprint"),
        ("S3 - Minor", "Sai lệch nhỏ, có cách vòng", "Medium", "Sprint sau"),
        ("S4 - Trivial", "Lỗi hiển thị, chính tả", "Low", "Khi có thời gian"),
    ], widths=[3, 7, 2.8, 3.6])
    r.p("Vòng đời bug trên Jira: Open > In Progress > Fixed (In Testing) > Retest > Closed; nếu kiểm tra lại vẫn lỗi "
        "thì Reopened và quay về In Progress. Mỗi bug ghi môi trường, bước tái hiện, kết quả mong đợi, kết quả thực tế, "
        "ảnh chụp và liên kết tới Story liên quan.", align=None)
    r.table(f"Bộ {len(D.TEST_CASES)} test case", ["Mã", "Tình huống", "Kết quả mong đợi", "Story", "Mức"],
            D.TEST_CASES, widths=[2.4, 5, 5.6, 1.6, 1.8], size=9.5)
    uat = [
        ("UAT-01", "Ứng viên đăng ký, Admin duyệt, Intern đăng nhập lần đầu bằng tài khoản được gửi qua email"),
        ("UAT-02", "Admin phân phòng ban và phân Mentor cho 6 Intern, kiểm tra giới hạn 5 Intern/Mentor"),
        ("UAT-03", "Mentor giao 3 task, Intern cập nhật tới REVIEW, Mentor duyệt 2 và trả lại 1"),
        ("UAT-04", "Intern chấm công 5 ngày có 1 ngày đi muộn, Mentor xem và xác nhận bảng công"),
        ("UAT-05", "Intern nộp 2 báo cáo tuần (1 đúng hạn, 1 muộn), Mentor phản hồi"),
        ("UAT-06", "Mentor đánh giá giữa kỳ và cuối kỳ, Intern xem điểm và xếp loại"),
        ("UAT-07", "Admin xem dashboard, lọc Intern chậm tiến độ, đối chiếu số liệu"),
        ("UAT-08", "Admin xuất kết quả Excel/PDF và kết thúc kỳ thực tập, kiểm tra dữ liệu bị khóa"),
    ]
    r.table("Kịch bản UAT", ["Mã", "Kịch bản"], uat, widths=[2, 14.4])
    rtm = defaultdict(list)
    for tc in D.TEST_CASES:
        rtm[tc[3]].append(tc[0])
    frs = defaultdict(list)
    for fr in D.FUNCTIONAL:
        frs[fr[3]].append(fr[0])
    r.table("Ma trận truy vết yêu cầu (RTM) tóm tắt", ["User Story", "Yêu cầu chức năng", "Test case"],
            [(us, ", ".join(frs[us]), ", ".join(rtm[us])) for us, *_ in D.USER_STORIES],
            widths=[2.4, 6, 8], size=10)


def ch9_control(r):
    r.h("9. Theo dõi và kiểm soát dự án", 1)
    r.p("Tiến độ được theo dõi trên Jira (board Scrum, Burndown chart, Velocity chart, Timeline) và đối chiếu với "
        "Gantt baseline. Chi phí và tiến độ tổng hợp được đo bằng phương pháp Giá trị thu được (Earned Value "
        "Management - EVM).")
    r.table("Các chỉ số EVM", ["Chỉ số", "Công thức", "Ý nghĩa"], [
        ("PV", "Giá trị công việc theo kế hoạch tới thời điểm đo", "Lẽ ra phải làm được bao nhiêu"),
        ("EV", "% hoàn thành thực tế x ngân sách công việc", "Thực tế làm được bao nhiêu"),
        ("AC", "Chi phí (giờ công quy đổi) thực tế đã bỏ ra", "Đã tốn bao nhiêu"),
        ("SV = EV - PV / SPI = EV / PV", "SPI < 1: chậm tiến độ", "Hiệu suất tiến độ"),
        ("CV = EV - AC / CPI = EV / AC", "CPI < 1: vượt chi phí", "Hiệu suất chi phí"),
        ("EAC = BAC / CPI", "Dự báo tổng chi phí khi hoàn thành", "Cần điều chỉnh gì"),
    ], widths=[4.6, 6.4, 5.4])
    when = date(2026, 10, 9)
    pv = planned_value(when)
    ev, ac = pv * 0.95, pv * 1.05
    bac = budget()[2]
    spi, cpi = ev / pv, ev / ac
    r.p(f"Ví dụ minh họa tại cuối tuần 6 ({when:%d/%m/%Y}, mốc M3). PV được tính từ kế hoạch: tổng giờ công đã lên "
        f"lịch tới ngày đó nhân đơn giá vai trò, PV = {money(pv)}; BAC = {money(bac)}. Giả định (số liệu minh họa, "
        "không phải thực tế) EV bằng 95% PV và AC bằng 105% PV:")
    r.table("Ví dụ tính EVM tại mốc M3 (số liệu minh họa)", ["Chỉ số", "Giá trị", "Nhận xét"], [
        ("PV", money(pv), "Tính từ kế hoạch"),
        ("EV (giả định)", money(ev), "95% khối lượng kế hoạch"),
        ("AC (giả định)", money(ac), "Tốn hơn kế hoạch 5%"),
        ("SV / SPI", f"{money(ev - pv)} / {spi:.2f}", "Chậm tiến độ khoảng 5%"),
        ("CV / CPI", f"{money(ev - ac)} / {cpi:.2f}", "Vượt chi phí khoảng 10%"),
        ("EAC = BAC / CPI", money(bac / cpi),
         f"Dự báo vượt {money(bac / cpi - bac)}, "
         + ("nằm trong quỹ dự phòng" if bac / cpi - bac <= bac * CONTINGENCY
            else f"vượt quỹ dự phòng {money(bac * CONTINGENCY)}: phải hành động ngay")),
    ], widths=[4, 5, 7.4])
    r.p("Hành động khi SPI hoặc CPI dưới 0,95: xem lại việc trên đường găng, dời Story ưu tiên thấp (Medium) sang "
        "Sprint sau theo quy trình Change Request, giảm thời gian ở các khung giờ PM để tăng giờ BE/FE, và dùng quỹ "
        "dự phòng khi cần.", align=None)
    r.table("Mẫu Weekly Status Report", ["Mục", "Nội dung cần ghi"], [
        ("Tuần / Sprint", "Tuần x, Sprint y, ngày báo cáo"),
        ("Tình trạng chung", "Xanh / Vàng / Đỏ kèm lý do"),
        ("Đã hoàn thành", "Danh sách work item Done (mã IM-xx), SP đạt"),
        ("Kế hoạch tuần tới", "Work item dự kiến, mốc sắp tới"),
        ("Chỉ số", "SP hoàn thành/cam kết, SPI, CPI, số bug mở theo mức"),
        ("Rủi ro và vấn đề", "Rủi ro mới hoặc thay đổi điểm, vấn đề cần hỗ trợ"),
        ("Change Request", "CR mới, CR đã duyệt/từ chối"),
    ], widths=[4, 12.4])


def ch10_closing(r):
    r.h("10. Kết thúc dự án", 1)
    r.p("Giai đoạn kết thúc (Sprint 5, tuần 12) xác nhận sản phẩm đáp ứng tiêu chí thành công, bàn giao đầy đủ và "
        "rút kinh nghiệm cho các dự án sau.")
    r.table("Danh mục bàn giao", ["Sản phẩm", "Định dạng", "Vai trò"], [
        ("Mã nguồn IMS v1.0.0 (tag trên GitHub)", "Repository + file zip", "BE, FE"),
        ("Tài liệu SRS, Use Case, User Story, RTM", "Word/PDF", "BA"),
        ("Tài liệu thiết kế: ERD, OpenAPI, wireframe", "PDF, Swagger, Figma", "BE, FE"),
        ("Test Plan, test case, Test Report, Postman collection", "Excel, JSON, PDF", "TEST"),
        ("Biên bản UAT Sign-off", "PDF có chữ ký", "BA"),
        ("User Manual 3 vai trò, Deployment Guide, Release Notes", "PDF, Markdown", "BA, BE"),
        ("Hồ sơ quản lý: Charter, WBS, Gantt, Risk Register, Status Report", "Word, Excel, Jira", "PM"),
        ("Báo cáo tổng kết, lessons learned", "Word", "PM"),
    ], widths=[8.6, 4.4, 3.4])
    r.p("Mẫu ghi nhận bài học kinh nghiệm (điền khi kết thúc dự án):", align=WD_ALIGN_PARAGRAPH.LEFT)
    r.table("Lessons learned", ["Hạng mục", "Điều làm tốt", "Điều cần cải thiện", "Hành động cho dự án sau"], [
        ("Lập kế hoạch", "", "", ""), ("Ước lượng", "", "", ""), ("Kiêm nhiệm 5 vai trò", "", "", ""),
        ("Chất lượng", "", "", ""), ("Công cụ Jira/GitHub", "", "", ""),
    ], widths=[3.8, 4.2, 4.2, 4.2])


def appendix(r):
    r.h("Phụ lục A. Triển khai trên Jira Cloud Free", 1)
    r.p("Toàn bộ backlog của dự án được sinh tự động từ cùng nguồn dữ liệu với báo cáo và Gantt, lưu trong file "
        "IMS_Jira_Import.csv: 10 Epic (tương ứng nhóm WBS 1-10), 16 Story (US-01..US-16, có Story Point và Acceptance "
        "Criteria), 74 Task và 31 Sub-task, gắn label theo vai trò, Sprint và mã WBS. Import vào project Scrum loại "
        "company-managed bằng Settings > System > External System Import > CSV để giữ quan hệ Epic - Story - Sub-task.")
    r.p("Các bước chi tiết (tạo site, tạo project, tạo 6 Sprint và lấy Sprint ID, sinh CSV, ánh xạ trường, kiểm tra "
        "sau import, cấu hình board cho một người làm 5 vai trò) được trình bày trong tài liệu riêng "
        "Huong_dan_Jira_IMS.docx đi kèm báo cáo này.")


def build(out):
    r = Report()
    cover(r)
    header_footer(r)
    for fn in (ch1_overview, ch2_charter, ch3_organization, ch4_requirements, ch5_design, ch6_plan,
               ch7_risk, ch8_test, ch9_control, ch10_closing, appendix):
        fn(r)
    r.doc.save(out)
    by_role, labour, subtotal, reserve, grand = budget()
    print(f"saved {out}: hours={total_hours()} budget={money(grand)} tables={len(r.doc.tables)}")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "IMS_Bao_cao_QLDAPM.docx")
